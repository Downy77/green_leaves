"""Turn uploaded sources into searchable text and bounded overlapping chunks."""

from __future__ import annotations

import re
from functools import lru_cache
from pathlib import Path

TEXT_SUFFIXES = {".txt", ".md", ".csv", ".tsv", ".json", ".py", ".js", ".ts", ".html", ".css", ".xml", ".yaml", ".yml", ".sql"}
MAX_EXTRACTED_CHARS = 300_000
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff"}


@lru_cache(maxsize=1)
def _ocr_engine():
    from rapidocr import RapidOCR

    return RapidOCR()


def _ocr_image(path: Path) -> str:
    result = _ocr_engine()(str(path))
    return "\n".join(result.txts or ())


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in TEXT_SUFFIXES:
        raw = path.read_bytes()
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = raw.decode("gb18030", errors="replace")
    elif suffix == ".pdf":
        from pypdf import PdfReader

        text = "\n\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages)
    elif suffix == ".docx":
        from docx import Document

        document = Document(str(path))
        text = "\n".join(
            [paragraph.text for paragraph in document.paragraphs]
            + ["\t".join(cell.text for cell in row.cells) for table in document.tables for row in table.rows]
        )
    elif suffix in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            lines = []
            length = 0
            for sheet in workbook.worksheets:
                lines.append(f"# {sheet.title}")
                for row in sheet.iter_rows(values_only=True):
                    values = [str(value).replace("\n", " ") if value is not None else "" for value in row]
                    if any(values):
                        line = "\t".join(values)
                        lines.append(line)
                        length += len(line)
                    if length >= MAX_EXTRACTED_CHARS:
                        break
                if length >= MAX_EXTRACTED_CHARS:
                    break
            text = "\n".join(lines)
        finally:
            workbook.close()
    elif suffix in IMAGE_SUFFIXES:
        text = _ocr_image(path)
    else:
        raise ValueError(f"暂不支持解析 {suffix or '该格式'} 文件；请使用文本、代码、CSV、PDF、DOCX、XLSX 或图片")
    text = text.replace("\x00", "").strip()
    if not text:
        raise ValueError("文件中没有可读取的文字内容；扫描版 PDF 需要先进行 OCR")
    return text[:MAX_EXTRACTED_CHARS]


def chunk_text(text: str, *, size: int = 700, overlap: int = 100) -> list[str]:
    """Prefer paragraph/sentence boundaries, retaining overlap for context."""
    cleaned = re.sub(r"[ \t]+", " ", text.replace("\r\n", "\n")).strip()
    if not cleaned:
        return []
    if size <= overlap or overlap < 0:
        raise ValueError("invalid chunk overlap")
    chunks: list[str] = []
    start = 0
    while start < len(cleaned):
        end = min(start + size, len(cleaned))
        if end < len(cleaned):
            boundary = max(cleaned.rfind(mark, start + size // 2, end) for mark in ("\n", "。", "！", "？", ".", "!", "?"))
            if boundary > start:
                end = boundary + 1
        piece = cleaned[start:end].strip()
        if piece:
            chunks.append(piece)
        if end >= len(cleaned):
            break
        start = max(start + 1, end - overlap)
    return chunks
