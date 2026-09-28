from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from fastapi import FastAPI
from fastapi.testclient import TestClient

from backend.core import db as database
from backend.routers import knowledge
from backend.services import knowledge_vectors, storage
from backend.services.knowledge_documents import chunk_text, extract_text


class KnowledgeTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory()
        root = Path(self.temp.name)
        self.uploads = root / "uploads"
        self.uploads.mkdir()
        self.patches = [
            patch.object(database, "DB_PATH", root / "test.db"),
            patch.object(storage, "UPLOAD_DIR", self.uploads),
            patch.object(knowledge, "UPLOAD_DIR", self.uploads),
            patch.object(knowledge, "enabled", return_value=True),
            patch.object(knowledge, "index_note", side_effect=self.fake_index),
            patch.object(knowledge, "delete_note"),
        ]
        self.indexed: dict[str, list[str]] = {}
        for item in self.patches:
            item.start()
        database.init_db()
        app = FastAPI()
        app.include_router(knowledge.router)
        self.client = TestClient(app)

    def tearDown(self) -> None:
        self.client.close()
        for item in reversed(self.patches):
            item.stop()
        self.temp.cleanup()

    def fake_index(self, note_id: str, content: str) -> int:
        pieces = chunk_text(content, size=80, overlap=10)
        self.indexed[note_id] = pieces
        return len(pieces)

    def test_upload_edit_search_relation_archive_and_download(self) -> None:
        response = self.client.post(
            "/api/knowledge/sources",
            files={"files": ("rules.txt", "客户跟进规则：三天内联系。".encode(), "text/plain")},
        )
        self.assertEqual(response.status_code, 200)
        note = response.json()["items"][0]
        note_id = note["id"]
        self.assertEqual(note["index_status"], "ready")
        self.assertEqual(self.client.get(f"/api/knowledge/notes/{note_id}/download").content, "客户跟进规则：三天内联系。".encode())
        self.assertIn("跟进规则", self.client.get(f"/api/knowledge/notes/{note_id}").json()["content"])
        self.assertEqual(len(self.indexed[note_id]), 1)

        second = self.client.post("/api/knowledge/sources/text", json={"title": "报价", "text": "报价要经主管确认"}).json()["id"]
        relation = self.client.post("/api/knowledge/relations", json={"source_id": note_id, "target_id": second}).json()
        self.assertTrue(relation["id"])
        self.assertEqual(len(self.client.get("/api/knowledge/graph").json()["edges"]), 1)

        with patch.object(knowledge_vectors, "enabled", return_value=True), patch.object(
            knowledge_vectors, "search", return_value=[{"knowledge_id": note_id, "chunk_index": 0, "content": "三天内联系客户", "score": 0.8}]
        ):
            results = self.client.get("/api/knowledge/search", params={"q": "何时回访客户"}).json()["items"]
        self.assertEqual(results[0]["note_id"], note_id)
        self.assertEqual(results[0]["match_excerpt"], "三天内联系客户")
        with patch.object(knowledge_vectors, "enabled", return_value=True), patch.object(
            knowledge_vectors, "search", return_value=[{"knowledge_id": note_id, "chunk_index": 0, "content": "不相关", "score": 0.3}]
        ):
            unrelated = self.client.get("/api/knowledge/search", params={"q": "量子物理"}).json()["items"]
        self.assertEqual(unrelated, [])

        changed = self.client.patch(f"/api/knowledge/notes/{note_id}", json={"content": "客户跟进规则：当天联系。"}).json()
        self.assertEqual(changed["index_status"], "ready")
        self.assertIn("当天", self.indexed[note_id][0])
        self.client.delete(f"/api/knowledge/notes/{note_id}")
        self.assertEqual(self.client.get(f"/api/knowledge/notes/{note_id}").status_code, 404)
        self.assertEqual(len(self.client.get("/api/knowledge/graph").json()["edges"]), 0)

    def test_unsupported_file_does_not_create_knowledge(self) -> None:
        result = self.client.post(
            "/api/knowledge/sources", files={"files": ("archive.bin", b"binary", "application/octet-stream")}
        ).json()
        self.assertFalse(result["ok"])
        self.assertEqual(result["items"], [])
        self.assertEqual(self.client.get("/api/knowledge/graph").json()["nodes"], [])

    def test_image_ocr_text_is_indexed(self) -> None:
        with patch("backend.services.knowledge_documents._ocr_image", return_value="图片里的客户守则"):
            result = self.client.post(
                "/api/knowledge/sources", files={"files": ("photo.png", b"image-data", "image/png")}
            ).json()
        self.assertEqual(result["items"][0]["index_status"], "ready")
        self.assertIn("客户守则", self.indexed[result["items"][0]["id"]][0])

    def test_chunk_overlap_and_text_extract(self) -> None:
        text = "第一段。" * 300
        chunks = chunk_text(text, size=90, overlap=15)
        self.assertGreater(len(chunks), 1)
        self.assertTrue(all(len(chunk) <= 90 for chunk in chunks))
        path = self.uploads / "note.md"
        path.write_text("# 知识\n内容", encoding="utf-8")
        self.assertIn("内容", extract_text(path))
        legacy = self.uploads / "legacy.csv"
        legacy.write_bytes("客户,报价\n小米,100".encode("gb18030"))
        self.assertIn("报价", extract_text(legacy))

    def test_office_document_and_sheet_text(self) -> None:
        from docx import Document
        from openpyxl import Workbook

        document = Document()
        document.add_paragraph("客户服务条款")
        doc_path = self.uploads / "terms.docx"
        document.save(doc_path)
        self.assertIn("客户服务条款", extract_text(doc_path))

        workbook = Workbook()
        workbook.active.title = "报价表"
        workbook.active.append(["客户", "价格"])
        workbook.active.append(["小米", 100])
        sheet_path = self.uploads / "prices.xlsx"
        workbook.save(sheet_path)
        extracted = extract_text(sheet_path)
        self.assertIn("报价表", extracted)
        self.assertIn("小米\t100", extracted)


if __name__ == "__main__":
    unittest.main()
